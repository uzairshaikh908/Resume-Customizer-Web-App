import re
import unicodedata
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Literal

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fpdf import FPDF
from pydantic import BaseModel, Field
from pypdf import PdfReader


MAX_FILE_SIZE = 5 * 1024 * 1024
MIN_JOB_DESCRIPTION_LENGTH = 50
MAX_JOB_DESCRIPTION_LENGTH = 15000
MAX_RESUME_TEXT_LENGTH = 30000


class AppError(Exception):

    def __init__(self, status_code: int, detail: str):
        super().__init__(detail)
        self.status_code = status_code
        self.detail = detail


SECTIONS = ("summary", "skills", "experience", "projects")
Section = Literal["summary", "skills", "experience", "projects"]


class Contact(BaseModel):
    name: str = Field(description="Candidate name copied from the resume, or empty.")
    email: str = Field(description="Email copied from the resume, or empty.")
    phone: str = Field(description="Phone copied from the resume, or empty.")
    location: str = Field(description="Location copied from the resume, or empty.")
    links: list[str] = Field(description="Profile or portfolio links copied from the resume.")


class RewrittenText(BaseModel):
    original: str = Field(description="The original text, copied verbatim from the resume.")
    customized: str = Field(description="The rewritten text. Must not add facts, numbers or skills.")


class RankedBullet(RewrittenText):
    priority: int = Field(description="Relevance to the job within this entry; 1 = most relevant.")


class Experience(BaseModel):
    title: str
    company: str
    location: str
    dates: str
    bullets: list[RankedBullet] = Field(description="Bullets in their original resume order.")


class Project(BaseModel):
    name: str
    bullets: list[RankedBullet] = Field(description="Bullets in their original resume order.")


class Education(BaseModel):
    degree: str
    institution: str
    dates: str
    details: str


class ExtraSection(BaseModel):
    title: str
    items: list[str]


class SkillLists(BaseModel):
    original: list[str] = Field(description="Every skill listed in the resume, verbatim.")
    customized: list[str] = Field(
        description="Only skills from the resume, reordered by relevance to the job."
    )


class Evidence(BaseModel):
    requirement: str
    evidence: str = Field(
        description="A short verbatim quote from the resume that supports it, or empty."
    )


class AIResume(BaseModel):
    job_title: str = Field(description="Job title from the job description, or empty.")
    company: str = Field(description="Company name from the job description, or empty.")
    contact: Contact
    headline: str = Field(description="Headline or current title copied from the resume, or empty.")
    summary: RewrittenText
    skills: SkillLists
    experience: list[Experience] = Field(description="Entries in their original resume order.")
    projects: list[Project]
    education: list[Education]
    certifications: list[str]
    other_sections: list[ExtraSection]
    job_skills: list[str] = Field(
        description="Skills, tools and technologies the job description asks for."
    )
    job_keywords: list[str] = Field(
        description="Other important terms from the job description (domains, methods, soft skills)."
    )
    responsibilities: list[Evidence] = Field(
        description="Main responsibilities from the job description with resume evidence."
    )
    suggested_keywords: list[str] = Field(
        description="Job keywords missing from the resume that the candidate could add only if true."
    )


class AICoverLetter(BaseModel):
    greeting: str
    paragraphs: list[str] = Field(description="Three or four short paragraphs.")
    closing: str
    signature: str = Field(description="Candidate name from the resume, or empty.")


class ExperienceEntry(BaseModel):
    title: str = ""
    company: str = ""
    location: str = ""
    dates: str = ""
    bullets: list[str] = []


class ProjectEntry(BaseModel):
    name: str = ""
    bullets: list[str] = []


class ResumeDocument(BaseModel):
    contact: Contact
    headline: str = ""
    summary: str = ""
    skills: list[str] = []
    experience: list[ExperienceEntry] = []
    projects: list[ProjectEntry] = []
    education: list[Education] = []
    certifications: list[str] = []
    other_sections: list[ExtraSection] = []


class CoverLetterDocument(BaseModel):
    greeting: str = ""
    paragraphs: list[str] = []
    closing: str = ""
    signature: str = ""


class ExportRequest(BaseModel):
    kind: Literal["resume", "cover_letter"]
    format: Literal["pdf", "docx"]
    resume: ResumeDocument | None = None
    cover_letter: CoverLetterDocument | None = None
    name: str = Field("", max_length=100)
    company: str = Field("", max_length=100)


ALLOWED_EXTENSIONS = {".pdf", ".docx"}

@dataclass
class ParsedResume:
    filename: str
    extension: str
    content: bytes
    text: str
    has_tables: bool = False

    @property
    def is_scanned_pdf(self) -> bool:
        return self.extension == ".pdf" and not self.text.strip()

def validate_job_description(job_description: str) -> str:
    job_description = (job_description or "").strip()

    if not job_description:
        raise AppError(400, "Please enter a job description.")

    if len(job_description) < MIN_JOB_DESCRIPTION_LENGTH:
        raise AppError(
            400,
            f"Please enter a job description with at least "
            f"{MIN_JOB_DESCRIPTION_LENGTH} characters.",
        )

    if len(job_description) > MAX_JOB_DESCRIPTION_LENGTH:
        raise AppError(
            400,
            f"The job description is too long. Please keep it under "
            f"{MAX_JOB_DESCRIPTION_LENGTH:,} characters.",
        )

    return job_description

def safe_filename(filename: str) -> str:
    name = Path(filename or "").name
    name = re.sub(r"[^A-Za-z0-9._ -]", "_", name).strip(" .")
    return name[:120] or "resume"

def _looks_like_docx(content: bytes) -> bool:
    if not content.startswith(b"PK\x03\x04"):
        return False
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            return "word/document.xml" in archive.namelist()
    except zipfile.BadZipFile:
        return False

def _extract_pdf(content: bytes) -> str:
    reader = PdfReader(BytesIO(content))

    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception:
            unlocked = 0
        if not unlocked:
            raise AppError(400, "Password-protected PDFs are not supported.")

    return "\n".join(page.extract_text() or "" for page in reader.pages)

def _extract_docx(content: bytes) -> tuple[str, bool]:
    document = Document(BytesIO(content))

    parts = [paragraph.text for paragraph in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))

    return "\n".join(parts), bool(document.tables)

def parse_resume(filename: str, content: bytes) -> ParsedResume:
    filename = safe_filename(filename)
    extension = Path(filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise AppError(400, "Only PDF and DOCX files are supported.")

    if not content:
        raise AppError(400, "The uploaded file is empty.")

    if len(content) > MAX_FILE_SIZE:
        raise AppError(413, "File size must not exceed 5 MB.")

    if extension == ".pdf" and not content.lstrip()[:5] == b"%PDF-":
        raise AppError(400, "This file is not a valid PDF.")
    if extension == ".docx" and not _looks_like_docx(content):
        raise AppError(400, "This file is not a valid DOCX document.")

    has_tables = False
    try:
        if extension == ".pdf":
            text = _extract_pdf(content)
        else:
            text, has_tables = _extract_docx(content)
    except AppError:
        raise
    except Exception:
        raise AppError(
            400,
            "Could not read the file. It may be corrupted. "
            "Please upload a valid PDF or DOCX.",
        )

    text = "\n".join(line.strip() for line in text.splitlines())
    text = re.sub(r"\n{3,}", "\n\n", text).strip()

    if not text and extension == ".docx":
        raise AppError(400, "No readable text was found in this document.")

    return ParsedResume(filename, extension, content, text, has_tables)

def parse_sections(raw: str) -> list[str]:
    sections = [s.strip().lower() for s in (raw or "").split(",") if s.strip()]
    invalid = [s for s in sections if s not in SECTIONS]
    if invalid:
        raise AppError(400, f"Unknown section(s): {', '.join(invalid)}. Allowed: {', '.join(SECTIONS)}.")
    if not sections:
        raise AppError(400, "Please select at least one section to customize.")
    return [s for s in SECTIONS if s in sections]



STANDARD_HEADINGS = {
    "Experience": ("experience", "work history", "employment", "professional experience"),
    "Education": ("education", "academic"),
    "Skills": ("skills", "technical skills", "core competencies"),
    "Summary": ("summary", "profile", "objective", "about me"),
}


def normalize(text: str) -> str:
    return " ".join(re.sub(r"[^a-z0-9+#]+", " ", (text or "").lower()).split())


class ResumeText:

    def __init__(self, text: str):
        self.normalized = normalize(text)
        self.padded = f" {self.normalized} "
        tokens = self.normalized.split()
        self.tokens = set(tokens)
        self.joined = {a + b for a, b in zip(tokens, tokens[1:])}
        self.joined |= {a + b + c for a, b, c in zip(tokens, tokens[1:], tokens[2:])}

    def has_term(self, term: str) -> bool:
        normalized = normalize(term)
        if not normalized:
            return False

        candidates = {normalized}
        stripped = re.sub(r"\s?js$", "", normalized).strip()
        if len(stripped) >= 3:
            candidates.add(stripped)

        for candidate in candidates:
            if f" {candidate} " in self.padded:
                return True
            if candidate.replace(" ", "") in self.tokens | self.joined:
                return True
        return False

    def supports_quote(self, quote: str, threshold: float = 0.7) -> bool:
        words = normalize(quote).split()
        if len(words) < 2:
            return False
        if " ".join(words) in self.normalized:
            return True
        return sum(word in self.tokens for word in words) / len(words) >= threshold


def numbers_in(text: str) -> set[str]:
    return set(re.findall(r"\d+(?:[.,]\d+)?", text or ""))


def dedupe(items: list[str]) -> list[str]:
    seen, result = set(), []
    for item in items:
        key = normalize(item)
        if key and key not in seen:
            seen.add(key)
            result.append(item.strip())
    return result


def _guard_bullets(
    bullets: list[RankedBullet], resume: ResumeText, customize: bool, warnings: list[str], where: str
) -> tuple[list[str], list[dict]]:
    kept = []
    for bullet in bullets:
        original = bullet.original.strip()
        if not original:
            continue
        if not resume.supports_quote(original):
            warnings.append(f"Removed a bullet in {where} that was not found in your resume.")
            continue

        customized = bullet.customized.strip() or original
        if not customize:
            customized = original
        elif not numbers_in(customized) <= numbers_in(original):
            warnings.append(
                f"Kept an original bullet in {where} because the rewrite added numbers "
                "that are not in your resume."
            )
            customized = original
        kept.append((bullet.priority, original, customized))

    if customize:
        kept.sort(key=lambda item: item[0])

    pairs = [{"original": o, "customized": c} for _, o, c in kept]
    return [c for _, _, c in kept], pairs


def build_resume(
    ai: AIResume, resume_text: str, sections: list[str]
) -> tuple[ResumeDocument, dict, dict, list[str]]:
    resume = ResumeText(resume_text)
    warnings: list[str] = []

    summary_original = ai.summary.original.strip()
    summary = ai.summary.customized.strip() if "summary" in sections else summary_original
    if "summary" in sections and not numbers_in(summary) <= numbers_in(resume_text):
        warnings.append(
            "Kept your original summary because the rewrite contained numbers "
            "that are not in your resume."
        )
        summary = summary_original

    original_skills = dedupe(ai.skills.original)
    proposed = dedupe(ai.skills.customized) if "skills" in sections else original_skills
    skills = [skill for skill in proposed if resume.has_term(skill)]
    unverified = [skill for skill in proposed if not resume.has_term(skill)]
    if unverified:
        warnings.append(
            "Removed skills not found in your resume: " + ", ".join(unverified) + "."
        )

    experience, experience_changes = [], []
    for entry in ai.experience:
        bullets, pairs = _guard_bullets(
            entry.bullets, resume, "experience" in sections, warnings,
            f"“{entry.title or entry.company}”",
        )
        experience.append(ExperienceEntry(
            title=entry.title, company=entry.company, location=entry.location,
            dates=entry.dates, bullets=bullets,
        ))
        experience_changes.append({"title": entry.title, "company": entry.company, "bullets": pairs})

    projects, project_changes = [], []
    for entry in ai.projects:
        bullets, pairs = _guard_bullets(
            entry.bullets, resume, "projects" in sections, warnings, f"“{entry.name}”"
        )
        projects.append(ProjectEntry(name=entry.name, bullets=bullets))
        project_changes.append({"name": entry.name, "bullets": pairs})

    document = ResumeDocument(
        contact=ai.contact,
        headline=ai.headline,
        summary=summary,
        skills=skills,
        experience=experience,
        projects=projects,
        education=ai.education,
        certifications=ai.certifications,
        other_sections=ai.other_sections,
    )

    changes = {
        "summary": {"original": summary_original, "customized": summary},
        "skills": {"original": original_skills, "customized": skills},
        "experience": experience_changes,
        "projects": project_changes,
    }

    job_skills = dedupe(ai.job_skills)
    skills_analysis = {
        "verified_existing": original_skills,
        "relevant_supported": [s for s in job_skills if resume.has_term(s)],
        "job_required_missing": [s for s in job_skills if not resume.has_term(s)],
        "suggested_keywords": [k for k in dedupe(ai.suggested_keywords) if not resume.has_term(k)],
    }

    return document, changes, skills_analysis, warnings


MATCH_WEIGHTS = {"skills": 0.5, "keywords": 0.3, "responsibilities": 0.2}


def match_score(ai: AIResume, resume_text: str) -> dict:
    resume = ResumeText(resume_text)

    job_skills = dedupe(ai.job_skills)
    skill_keys = {normalize(s) for s in job_skills}
    keywords = [k for k in dedupe(ai.job_keywords) if normalize(k) not in skill_keys]
    responsibilities = [r for r in ai.responsibilities if r.requirement.strip()]

    groups = {
        "skills": [(s, resume.has_term(s)) for s in job_skills],
        "keywords": [(k, resume.has_term(k)) for k in keywords],
        "responsibilities": [
            (r.requirement, bool(r.evidence) and resume.supports_quote(r.evidence))
            for r in responsibilities
        ],
    }
    evidence = {
        r.requirement: r.evidence
        for r in responsibilities
        if r.evidence and resume.supports_quote(r.evidence)
    }

    components, weighted, total_weight = [], 0.0, 0.0
    for name, items in groups.items():
        matched = [item for item, ok in items if ok]
        missing = [item for item, ok in items if not ok]
        ratio = len(matched) / len(items) if items else None
        if ratio is not None:
            weighted += MATCH_WEIGHTS[name] * ratio
            total_weight += MATCH_WEIGHTS[name]
        components.append({
            "name": name,
            "weight": MATCH_WEIGHTS[name],
            "matched": matched,
            "missing": missing,
            "score": None if ratio is None else round(ratio * 100),
        })

    return {
        "score": round(100 * weighted / total_weight) if total_weight else 0,
        "components": components,
        "evidence": evidence,
        "method": (
            "Skills 50%, keywords 30%, responsibilities 20%: the share of each found in "
            "your original resume (skills and keywords by exact term match, responsibilities "
            "by a supporting quote that exists in your resume). Groups the job description "
            "does not mention are left out. This is an estimate, not a real ATS score."
        ),
    }


def ats_analysis(resume_text: str, is_scanned: bool, has_tables: bool, job_terms: list[str]) -> dict:
    resume = ResumeText(resume_text)
    words = resume_text.split()
    lowered_lines = [line.strip().lower().rstrip(":") for line in resume_text.splitlines()]

    found_headings = [
        heading for heading, variants in STANDARD_HEADINGS.items()
        if any(line.startswith(variants) and len(line) <= 40 for line in lowered_lines)
    ]
    terms = dedupe(job_terms)
    matched_terms = [t for t in terms if resume.has_term(t)]
    keyword_ratio = len(matched_terms) / len(terms) if terms else 1.0
    has_phone = any(
        10 <= len(re.sub(r"\D", "", candidate)) <= 15
        for candidate in re.findall(r"\+?\d[\d\s().-]{7,}\d", resume_text)
    )
    has_bullets = bool(re.search(r"^\s*([-•*●▪◦]|\d+\.)\s+", resume_text, re.MULTILINE))

    checks = [
        {
            "name": "Machine-readable text",
            "weight": 20,
            "ratio": 0.0 if is_scanned else 1.0,
            "detail": "Scanned / image-only PDF" if is_scanned else "Text can be extracted",
            "tip": "Export your resume as a text-based PDF or DOCX so ATS software can read it.",
        },
        {
            "name": "Email address",
            "weight": 10,
            "ratio": 1.0 if re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", resume_text) else 0.0,
            "detail": "Found" if re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", resume_text) else "Not found",
            "tip": "Add a professional email address near the top.",
        },
        {
            "name": "Phone number",
            "weight": 5,
            "ratio": 1.0 if has_phone else 0.0,
            "detail": "Found" if has_phone else "Not found",
            "tip": "Add a phone number to your contact details.",
        },
        {
            "name": "Standard section headings",
            "weight": 20,
            "ratio": len(found_headings) / len(STANDARD_HEADINGS),
            "detail": "Found: " + (", ".join(found_headings) or "none"),
            "tip": "Use standard headings such as Summary, Experience, Education and Skills.",
        },
        {
            "name": "Job keyword coverage",
            "weight": 25,
            "ratio": min(1.0, keyword_ratio / 0.6),
            "detail": f"{len(matched_terms)} of {len(terms)} job terms appear in your resume",
            "tip": "Mention the job's key skills and terms you genuinely have, using the same wording.",
        },
        {
            "name": "Length",
            "weight": 10,
            "ratio": 1.0 if 250 <= len(words) <= 1200 else 0.5 if 120 <= len(words) <= 1600 else 0.0,
            "detail": f"{len(words)} words",
            "tip": "Aim for roughly 250–1,200 words (one to two pages).",
        },
        {
            "name": "Bullet points",
            "weight": 5,
            "ratio": 1.0 if has_bullets else 0.0,
            "detail": "Used" if has_bullets else "Not detected",
            "tip": "Describe achievements as short bullet points.",
        },
        {
            "name": "Simple layout (no tables)",
            "weight": 5,
            "ratio": 0.0 if has_tables else 1.0,
            "detail": "Tables found" if has_tables else "No tables detected",
            "tip": "Avoid tables and multi-column layouts; some ATS parsers read them out of order.",
        },
    ]

    score = sum(c["weight"] * c["ratio"] for c in checks) / sum(c["weight"] for c in checks)
    for check in checks:
        check["passed"] = check["ratio"] >= 0.99

    return {
        "score": round(score * 100),
        "checks": checks,
        "improvements": [c["tip"] for c in checks if not c["passed"]],
        "method": (
            "Weighted checks of common ATS-readiness signals in your original resume. "
            "This is an estimate; it does not reproduce any specific ATS product."
        ),
    }


MEDIA_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

_PDF_REPLACEMENTS = {
    "•": "-", "●": "-", "▪": "-", "–": "-", "—": "-",
    "‘": "'", "’": "'", "“": '"', "”": '"', "…": "...",
    " ": " ", "→": "->",
}


def export_filename(kind: str, fmt: str, name: str, company: str) -> str:
    def part(text: str) -> str:
        ascii_text = unicodedata.normalize("NFKD", text or "").encode("ascii", "ignore").decode()
        return re.sub(r"[^A-Za-z0-9]+", "_", ascii_text).strip("_")[:40]

    label = "Cover_Letter" if kind == "cover_letter" else "Resume"
    pieces = [part(name), label, part(company)]
    return "_".join(p for p in pieces if p) + f".{fmt}"


def _contact_line(resume: ResumeDocument) -> str:
    c = resume.contact
    return " | ".join(p for p in [c.email, c.phone, c.location, *c.links] if p)


def _resume_sections(resume: ResumeDocument) -> list[tuple[str, list]]:
    sections = []
    if resume.summary:
        sections.append(("Professional Summary", [("text", resume.summary)]))
    if resume.skills:
        sections.append(("Skills", [("text", ", ".join(resume.skills))]))
    if resume.experience:
        blocks = []
        for job in resume.experience:
            title = " - ".join(p for p in [job.title, job.company] if p)
            blocks.append(("entry", title, " | ".join(p for p in [job.location, job.dates] if p)))
            blocks += [("bullet", b) for b in job.bullets]
        sections.append(("Experience", blocks))
    if resume.projects:
        blocks = []
        for project in resume.projects:
            blocks.append(("entry", project.name, ""))
            blocks += [("bullet", b) for b in project.bullets]
        sections.append(("Projects", blocks))
    if resume.education:
        blocks = []
        for edu in resume.education:
            title = " - ".join(p for p in [edu.degree, edu.institution] if p)
            blocks.append(("entry", title, edu.dates))
            if edu.details:
                blocks.append(("text", edu.details))
        sections.append(("Education", blocks))
    if resume.certifications:
        sections.append(("Certifications", [("bullet", c) for c in resume.certifications]))
    for extra in resume.other_sections:
        if extra.title and extra.items:
            sections.append((extra.title, [("bullet", item) for item in extra.items]))
    return sections


def _new_docx() -> Document:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for section in document.sections:
        section.left_margin = section.right_margin = Pt(54)
        section.top_margin = section.bottom_margin = Pt(54)
    return document


def _docx_bytes(document: Document) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def resume_docx(resume: ResumeDocument) -> bytes:
    document = _new_docx()

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run(resume.contact.name or "Resume")
    run.bold = True
    run.font.size = Pt(20)

    for line in [resume.headline, _contact_line(resume)]:
        if line:
            paragraph = document.add_paragraph(line)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for heading, blocks in _resume_sections(resume):
        document.add_heading(heading, level=2)
        for block in blocks:
            if block[0] == "text":
                document.add_paragraph(block[1])
            elif block[0] == "bullet":
                document.add_paragraph(block[1], style="List Bullet")
            else:
                paragraph = document.add_paragraph()
                paragraph.add_run(block[1]).bold = True
                if block[2]:
                    paragraph.add_run(f"  |  {block[2]}")

    return _docx_bytes(document)


def cover_letter_docx(letter: CoverLetterDocument) -> bytes:
    document = _new_docx()
    for text in [letter.greeting, *letter.paragraphs, letter.closing, letter.signature]:
        if text:
            document.add_paragraph(text)
    return _docx_bytes(document)


def _latin1(text: str) -> str:
    for char, replacement in _PDF_REPLACEMENTS.items():
        text = text.replace(char, replacement)
    if not all(ord(ch) < 256 for ch in text):
        text = unicodedata.normalize("NFKD", text)
        text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return text.encode("latin-1", "replace").decode("latin-1")


class _PDF(FPDF):
    def write_line(self, text: str, size: float = 10.5, style: str = "", align: str = "L", gap: float = 1.5):
        self.set_font("Helvetica", style, size)
        self.multi_cell(0, size * 0.5, _latin1(text), align=align, new_x="LMARGIN", new_y="NEXT")
        self.ln(gap)

    def bullet(self, text: str):
        self.set_font("Helvetica", "", 10.5)
        self.cell(5, 5.25, "-")
        self.multi_cell(0, 5.25, _latin1(text), new_x="LMARGIN", new_y="NEXT")
        self.ln(1)

    def heading(self, text: str):
        self.ln(2)
        self.set_font("Helvetica", "B", 12)
        self.cell(0, 7, _latin1(text.upper()), new_x="LMARGIN", new_y="NEXT")
        y = self.get_y()
        self.line(self.l_margin, y, self.w - self.r_margin, y)
        self.ln(2)


def _new_pdf() -> _PDF:
    pdf = _PDF(format="A4")
    pdf.set_margins(18, 18, 18)
    pdf.set_auto_page_break(True, margin=18)
    pdf.add_page()
    return pdf


def resume_pdf(resume: ResumeDocument) -> bytes:
    pdf = _new_pdf()
    pdf.write_line(resume.contact.name or "Resume", size=20, style="B", align="C", gap=2)
    for line in [resume.headline, _contact_line(resume)]:
        if line:
            pdf.write_line(line, align="C", gap=1)

    for heading, blocks in _resume_sections(resume):
        pdf.heading(heading)
        for block in blocks:
            if block[0] == "text":
                pdf.write_line(block[1])
            elif block[0] == "bullet":
                pdf.bullet(block[1])
            else:
                title = block[1] + (f"  |  {block[2]}" if block[2] else "")
                pdf.write_line(title, style="B", gap=1)

    return bytes(pdf.output())


def cover_letter_pdf(letter: CoverLetterDocument) -> bytes:
    pdf = _new_pdf()
    for text in [letter.greeting, *letter.paragraphs, letter.closing, letter.signature]:
        if text:
            pdf.write_line(text, size=11, gap=4)
    return bytes(pdf.output())
